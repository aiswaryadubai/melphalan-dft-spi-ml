"""Reproduce SPI rankings, weight sensitivity, and ML validation.

This script intentionally preserves the signed-value normalization used in the
current ranking tables. See README.md before reusing or modifying the scores.
"""

from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sklearn.ensemble import GradientBoostingRegressor, RandomForestRegressor
from sklearn.impute import SimpleImputer
from sklearn.metrics import mean_absolute_error, mean_squared_error, r2_score
from sklearn.model_selection import KFold, cross_validate, train_test_split
from sklearn.pipeline import Pipeline


ROOT = Path(__file__).resolve().parents[1]
DATA_FILE = ROOT / "data" / "melphalan_76_systems.xlsx"
RESULTS = ROOT / "results"
RANDOM_STATE = 42

FEATURES = [
    "Distance_A",
    "HOMO_eV",
    "LUMO_eV",
    "Bandgap_eV",
    "Mulliken_Charge_Dopant",
    "Charge_Transfer_e",
    "E_ads_DFT_eV",
    "Recoverytime_sec",
    "electronegativity_eV",
    "chemical potential_eV",
    "hardness_eV",
    "electrophilicity index_eV",
    "Dopant_Count",
    "Total Atoms",
]

SCENARIOS = {
    "Sensor baseline": ("sensor", 0.40, 0.40, 0.20, 0.00),
    "Sensor equal": ("sensor", 1 / 3, 1 / 3, 1 / 3, 0.00),
    "Sensor adsorption dominant": ("sensor", 0.50, 0.30, 0.20, 0.00),
    "Sensor charge-transfer dominant": ("sensor", 0.30, 0.50, 0.20, 0.00),
    "Sensor recovery dominant": ("sensor", 0.30, 0.30, 0.40, 0.00),
    "Device baseline": ("device", 0.30, 0.30, 0.20, 0.20),
    "Device equal": ("device", 0.25, 0.25, 0.25, 0.25),
    "Device electronic-response dominant": ("device", 0.25, 0.25, 0.15, 0.35),
    "Device recovery dominant": ("device", 0.25, 0.25, 0.35, 0.15),
}


def minmax(series: pd.Series) -> pd.Series:
    """Return Min–Max scaled values, or zeros for a constant column."""
    numeric = pd.to_numeric(series, errors="coerce")
    span = numeric.max() - numeric.min()
    if span == 0 or pd.isna(span):
        return pd.Series(0.0, index=series.index)
    return (numeric - numeric.min()) / span


def load_data() -> pd.DataFrame:
    data = pd.read_excel(DATA_FILE)
    data.columns = data.columns.str.strip()
    data = data.rename(columns={"Adosption_Site": "Adsorption_Site"})
    if len(data) != 76:
        raise ValueError(f"Expected 76 systems, found {len(data)}")
    return data


def calculate_components(data: pd.DataFrame) -> pd.DataFrame:
    out = data.copy()
    out["Eads_score"] = minmax(out["E_ads_DFT_eV"])
    out["Q_score"] = minmax(out["Charge_Transfer_e"])
    out["Recovery_score"] = 1.0 - minmax(out["Recoverytime_sec"])
    out["Gap_score"] = 1.0 - minmax(out["Bandgap_eV"])
    return out


def calculate_rankings(data: pd.DataFrame) -> pd.DataFrame:
    out = data.copy()
    out["SPI_sensor"] = (
        0.40 * out["Eads_score"]
        + 0.40 * out["Q_score"]
        + 0.20 * out["Recovery_score"]
    )
    out["SPI_device"] = (
        0.30 * out["Eads_score"]
        + 0.30 * out["Q_score"]
        + 0.20 * out["Recovery_score"]
        + 0.20 * out["Gap_score"]
    )
    out["Rank_SPI_sensor"] = out["SPI_sensor"].rank(
        ascending=False, method="min"
    ).astype(int)
    out["Rank_SPI_device"] = out["SPI_device"].rank(
        ascending=False, method="min"
    ).astype(int)
    out["Rank_shift"] = out["Rank_SPI_sensor"] - out["Rank_SPI_device"]
    return out


def weight_sensitivity(data: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    detail = data[["System_ID", "Nanostructure_Type", "Dopant", "Medium"]].copy()
    summary_rows = []
    baseline_ranks = {}

    for name, (index_type, w_e, w_q, w_r, w_g) in SCENARIOS.items():
        score = (
            w_e * data["Eads_score"]
            + w_q * data["Q_score"]
            + w_r * data["Recovery_score"]
            + w_g * data["Gap_score"]
        )
        rank = score.rank(ascending=False, method="min").astype(int)
        detail[f"{name} score"] = score
        detail[f"{name} rank"] = rank
        if name in {"Sensor baseline", "Device baseline"}:
            baseline_ranks[index_type] = rank

    for name, (index_type, w_e, w_q, w_r, w_g) in SCENARIOS.items():
        score = detail[f"{name} score"]
        rank = detail[f"{name} rank"]
        top_index = score.idxmax()
        ni_mask = data["System_ID"].eq("BN_Ni1")
        ordered = score.sort_values(ascending=False)
        summary_rows.append(
            {
                "Scenario": name,
                "Index": index_type,
                "w(Eads)": w_e,
                "w(Q)": w_q,
                "w(recovery)": w_r,
                "w(gap)": w_g,
                "BNNS-Ni score": float(score.loc[ni_mask].iloc[0]),
                "BNNS-Ni rank": int(rank.loc[ni_mask].iloc[0]),
                "Top-ranked system": data.loc[top_index, "Nanostructure_Type"],
                "Top System_ID": data.loc[top_index, "System_ID"],
                "Medium": data.loc[top_index, "Medium"],
                "Spearman rho vs baseline": rank.corr(
                    baseline_ranks[index_type], method="spearman"
                ),
                "Lead over Rank 2": float(ordered.iloc[0] - ordered.iloc[1]),
            }
        )
    return pd.DataFrame(summary_rows), detail


def model_pipeline(model):
    return Pipeline(
        [("imputer", SimpleImputer(strategy="median")), ("model", model)]
    )


def machine_learning(data: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    X = data[FEATURES].apply(pd.to_numeric, errors="coerce")
    X = X.replace([np.inf, -np.inf], np.nan)
    y = pd.to_numeric(data["SPI_sensor"], errors="coerce")
    valid = y.notna()
    X, y = X.loc[valid].reset_index(drop=True), y.loc[valid].reset_index(drop=True)

    X_train, X_test, y_train, y_test = train_test_split(
        X, y, test_size=0.20, random_state=RANDOM_STATE
    )
    models = {
        "Random Forest": RandomForestRegressor(
            n_estimators=500, random_state=RANDOM_STATE
        ),
        "Gradient Boosting": GradientBoostingRegressor(random_state=RANDOM_STATE),
    }
    try:
        from xgboost import XGBRegressor

        models["XGBoost"] = XGBRegressor(
            n_estimators=300,
            learning_rate=0.03,
            max_depth=3,
            subsample=0.8,
            colsample_bytree=0.8,
            random_state=RANDOM_STATE,
            objective="reg:squarederror",
        )
    except ImportError:
        print("XGBoost is not installed; continuing with scikit-learn models.")

    rows = []
    for name, model in models.items():
        pipeline = model_pipeline(model)
        pipeline.fit(X_train, y_train)
        prediction = pipeline.predict(X_test)
        rows.append(
            {
                "Model": name,
                "Hold-out R2": r2_score(y_test, prediction),
                "Hold-out MAE": mean_absolute_error(y_test, prediction),
                "Hold-out RMSE": np.sqrt(mean_squared_error(y_test, prediction)),
            }
        )

    cv = KFold(n_splits=5, shuffle=True, random_state=RANDOM_STATE)
    gb = model_pipeline(GradientBoostingRegressor(random_state=RANDOM_STATE))
    cv_results = cross_validate(
        gb,
        X,
        y,
        cv=cv,
        scoring={
            "R2": "r2",
            "MAE": "neg_mean_absolute_error",
            "MSE": "neg_mean_squared_error",
        },
    )
    fold_table = pd.DataFrame(
        {
            "Fold": np.arange(1, 6),
            "R2": cv_results["test_R2"],
            "MAE": -cv_results["test_MAE"],
            "RMSE": np.sqrt(-cv_results["test_MSE"]),
        }
    )
    return pd.DataFrame(rows), fold_table


def plot_cv(folds: pd.DataFrame) -> None:
    mean_r2 = folds["R2"].mean()
    fig, ax = plt.subplots(figsize=(7, 5))
    bars = ax.bar(
        [f"Fold {fold}" for fold in folds["Fold"]],
        folds["R2"],
        color="#2878B5",
        edgecolor="black",
        linewidth=0.7,
    )
    ax.axhline(
        mean_r2,
        color="#D62728",
        linestyle="--",
        linewidth=1.5,
        label=f"Mean R² = {mean_r2:.4f}",
    )
    for bar, score in zip(bars, folds["R2"]):
        ax.text(
            bar.get_x() + bar.get_width() / 2,
            score + 0.02,
            f"{score:.3f}",
            ha="center",
            va="bottom",
        )
    ax.set_ylim(min(0, folds["R2"].min() - 0.1), 1.10)
    ax.set_xlabel("Validation fold")
    ax.set_ylabel("Cross-validation R²")
    ax.set_title("Five-Fold Cross-Validation of Gradient Boosting")
    ax.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(RESULTS / "Gradient_Boosting_5Fold_CV.png", dpi=600)
    plt.close(fig)


def format_workbook(path: Path) -> None:
    """Apply lightweight publication-ready formatting to generated tables."""
    workbook = load_workbook(path)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[1].height = 30
        for column_cells in sheet.columns:
            letter = column_cells[0].column_letter
            longest = max(
                len(str(cell.value)) if cell.value is not None else 0
                for cell in column_cells[:100]
            )
            sheet.column_dimensions[letter].width = min(max(longest + 2, 11), 28)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="center")
    workbook.save(path)


def main() -> None:
    RESULTS.mkdir(exist_ok=True)
    source = load_data()
    scored = calculate_components(source)
    rankings = calculate_rankings(scored)
    sensitivity_summary, sensitivity_detail = weight_sensitivity(scored)
    model_results, folds = machine_learning(rankings)

    ranking_file = RESULTS / "complete_rankings.xlsx"
    sensitivity_summary_file = RESULTS / "weight_sensitivity_summary.xlsx"
    sensitivity_detail_file = RESULTS / "weight_sensitivity_all_systems.xlsx"
    rankings.sort_values("Rank_SPI_sensor").to_excel(ranking_file, index=False)
    sensitivity_summary.to_excel(
        sensitivity_summary_file, index=False
    )
    sensitivity_detail.to_excel(
        sensitivity_detail_file, index=False
    )
    for workbook_path in [ranking_file, sensitivity_summary_file, sensitivity_detail_file]:
        format_workbook(workbook_path)
    model_results.to_csv(RESULTS / "model_performance.csv", index=False)
    folds.to_csv(RESULTS / "five_fold_cross_validation.csv", index=False)
    plot_cv(folds)

    ni = rankings.loc[rankings["System_ID"].eq("BN_Ni1")].iloc[0]
    if int(ni["Rank_SPI_sensor"]) != 1 or int(ni["Rank_SPI_device"]) != 1:
        raise AssertionError("Preserved BNNS–Ni baseline ranking was not reproduced.")

    print("Analysis completed for", len(rankings), "systems")
    print("BNNS–Ni SPI_sensor:", f"{ni['SPI_sensor']:.10f}")
    print("BNNS–Ni SPI_device:", f"{ni['SPI_device']:.10f}")
    print("Mean five-fold R²:", f"{folds['R2'].mean():.4f}")
    print("Mean five-fold MAE:", f"{folds['MAE'].mean():.4f}")
    print("Mean five-fold RMSE:", f"{folds['RMSE'].mean():.4f}")


if __name__ == "__main__":
    main()
